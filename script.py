import os
import pandas as pd
from docx import Document
import win32com.client  # For Word-to-PDF conversion

# Paths to input files and output directory
excel_path = os.path.abspath("data.xlsx")
word_template_path = os.path.abspath("template.docx")
output_dir = os.path.abspath("output_pdfs")
os.makedirs(output_dir, exist_ok=True)

# Read data from Excel
data = pd.read_excel(excel_path)
data.columns = data.columns.str.strip().str.lower()


# Replace placeholders in the Word document
def replace_placeholders(doc, replacements):
    for para in doc.paragraphs:
        for placeholder, replacement in replacements.items():
            if placeholder in para.text:
                for run in para.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, replacement)
                        run.bold = True


# Convert Word to PDF
def convert_docx_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(docx_path)
    doc.SaveAs(pdf_path, FileFormat=17)
    doc.Close()
    word.Quit()


# Process each row in the Excel sheet
def process_row(row, company_folder, invitation_number):
    doc = Document(word_template_path)

    name_lower = row["name"].strip().lower()
    if "mr" in name_lower:
        salutation = "Dear Sir,"
    elif "ms" in name_lower:
        salutation = "Dear Ma'am,"
    else:
        salutation = "Dear Sir/Ma'am"

    replacements = {
        "{name}": row["name"],
        "{company}": row["company"],
        "{post}": row["post"],
        "{s}": salutation,
    }
    replace_placeholders(doc, replacements)

    subfolder_path = os.path.join(company_folder, str(invitation_number))
    os.makedirs(subfolder_path, exist_ok=True)

    temp_word_path = os.path.join(subfolder_path, "temp.docx")
    pdf_path = os.path.join(subfolder_path, f"{row['company']}_Invitation_MMMUT.pdf")
    doc.save(temp_word_path)
    convert_docx_to_pdf(temp_word_path, pdf_path)
    os.remove(temp_word_path)


# Process all rows and create structured folders
company_folders = {}

for index, row in data.iterrows():
    company_name = row["company"].strip()

    if company_name not in company_folders:
        company_folder = os.path.join(output_dir, company_name)
        os.makedirs(company_folder, exist_ok=True)
        company_folders[company_name] = {"folder": company_folder, "count": 1}

    process_row(row, company_folders[company_name]["folder"], company_folders[company_name]["count"])
    company_folders[company_name]["count"] += 1

# Create HR Details.xlsx in each company folder
for company_name, info in company_folders.items():
    company_data = data[data["company"].str.strip() == company_name]
    hr_details_path = os.path.join(info["folder"], "HR Details.xlsx")
    company_data.to_excel(hr_details_path, index=False)
